#!/usr/bin/env python3
"""Move the 2026-09-30 role-polish cards to DEV_COMPLETED with their dated
evidence blocks (kanban-work skill §7). Terse by directive: evidence, no prose.

Run: python3 scripts/kanban-cards-20260930-role-ui-polish-devcompleted.py
"""
import pathlib

import docx

REPO = pathlib.Path('/Volumes/LexarSSD/projects/silversea-prod')
KANBAN = pathlib.Path(REPO.joinpath('.kanban-dir').read_text().strip())
DATE = '2026-09-30'

# card file -> (Trạng thái line, DEV COMPLETED block lines, verify steps or None)
DONE = {
    '20260930_213-ops-wallet-expense-history-404.docx': (
        'Trạng thái: DEV COMPLETED — chờ deploy staging + QA',
        [
            '[2026-09-30] DEV COMPLETED — AC1/AC2/AC3: `/ops/wallet` hết dòng lỗi; '
            '`GET /api/ops/wallet/expenses` → 200 với 8 dòng thật của giaonhan; mục "Lịch sử chi phí" '
            'render ở 390/768/1440, 0 lỗi console/API. AC4: test đỏ-trước rồi xanh '
            '(`qa/2026-09-30_card213_regression-red.log` → `-green.log`). AC5: `tsc --noEmit` 0. '
            'Dọn dữ liệu: 6 link mồ côi đã xoá ở DB local (6 dòng DELETE).',
            'Chưa làm (có thẻ riêng): 28 link DRIVER còn "join-broken" (dòng nguồn còn, chỉ join '
            'trips/drivers hỏng) — không xoá, vì nguồn vẫn tồn tại.',
        ],
        [
            'Mở http://localhost:7175/ops/wallet, đăng nhập giaonhan / Abc123.',
            'Kéo xuống mục "Lịch sử chi phí" — phải có bảng dòng chi phí, KHÔNG có dòng '
            '"Không tải được lịch sử chi phí".',
            'Mở DevTools → Network, tìm `wallet/expenses` — phải 200.',
        ],
    ),
    '20260930_215-dispatch-pill-badge-and-unassigned-dup.docx': (
        'Trạng thái: DEV COMPLETED — chờ deploy staging + QA',
        [
            '[2026-09-30] DEV COMPLETED — AC1: chip trạng thái/chiều là text thường + chấm màu; '
            'đo computed style trên /dispatch-detail 390: không còn phần tử `border-radius: 9999px` '
            '(trước: 9999px, padding 2px 8px). AC2: bản ghi chưa điều xe còn ĐÚNG một tín hiệu '
            '"chưa" — đếm chuỗi "Chưa điều xe" = 0 ở 390 và 1440 (trước: 1 mỗi bản ghi). AC3: '
            'ảnh 390/1440 kèm theo; 4 design-lock `no-pill-in-grid` mới giữ 0 pill trong cả 2 lưới.',
            'Cổng: frontend `tsc -b` 0; vitest `src/features/dispatch` 39 file / 362 test xanh; '
            '`pnpm design:lock` 200/201 (1 đỏ ở /customers w768, không thuộc thẻ này).',
        ],
        [
            'Mở /dispatch-detail (dieuvan / Abc123) ở 390px.',
            'Mỗi bản ghi: ô "Điều phối" chỉ còn "Chưa phân xe" — không còn dòng "Chưa điều xe".',
            'Khi có bản ghi đã phân xe: trạng thái hiện dạng chữ + chấm màu, không có bong bóng bo tròn.',
        ],
    ),
    '20260930_216-filler-copy-sweep-four-roles.docx': (
        'Trạng thái: DEV COMPLETED — chờ deploy staging + QA',
        [
            '[2026-09-30] DEV COMPLETED — AC1: 14 câu dài được rút, 2 khối bị xoá, trên 12 file '
            'thuộc bề mặt 4 vai trò (ví dụ `/my-earnings`: 211 → 73 ký tự; `/ops/wallet` bỏ hẳn dòng '
            '"Chỉ hiển thị các khoản tạm ứng/hoàn ứng của bạn; sổ chỉ đọc."). AC2: giữ nguyên mọi '
            'thông điệp lỗi/quyền/hành động phá huỷ. AC3: ảnh + text dump 390/1440 trong qa/.',
            'Cổng: vitest trên các cây bị chạm 379 test (1 đỏ `OpsExpenseFormModal.containers` A1 là '
            'timeout 5,3s/5s có sẵn — xanh khi chạy lại; lớp flake đã có thẻ 20260928_185).',
            'Chưa làm: 282 ứng viên còn lại ngoài 4 vai trò (kế toán, users, config) — nêu trong thẻ.',
        ],
        [
            'Mở /my-earnings (laixe) ở 390px — khối "Ghi chú" chỉ còn MỘT câu ngắn.',
            'Mở /ops/wallet (giaonhan) — mục "Sổ quỹ" không còn dòng giải thích dài phía trên bảng.',
            'Mở form tạo lô (/shipments/new → khách hàng mới) — hộp thoại chỉ còn câu ngắn về công nợ.',
        ],
    ),
    '20260930_217-role-sweep-tapfloor-40-not-44.docx': (
        'Trạng thái: DEV COMPLETED — chờ deploy staging + QA',
        [
            '[2026-09-30] DEV COMPLETED — AC1: `role-ui-sweep.mjs` + `mobile-ux-sweep.mjs` dùng sàn '
            '40px (trước 44/43.5); `role-ui-sweep-report.mjs` in "TAP<40". AC2: ngưỡng ghi rõ trong '
            'chú thích + log. AC3: sweep lại 390/768/1440 cho 4 vai trò → mục "tap targets <40" '
            'TRỐNG (trước: 24 dòng TAP<44, trong đó 0 control thật sự dưới 40px).',
            'Bằng chứng: qa/2026-09-30_role-sweep_report-after-tapfloor-fix.txt, '
            'qa/2026-09-30_role-sweep_sweep.log, báo cáo trước khi sửa giữ ở '
            'qa/2026-09-30_role-sweep_before-tapfloor-fix.json.',
        ],
        None,
    ),
    '20260930_218-ops-wallet-density.docx': (
        'Trạng thái: DEV COMPLETED — chờ deploy staging + QA',
        [
            '[2026-09-30] DEV COMPLETED — AC1: ba số dư nằm trong MỘT dải `SummaryRail` ở 390/768/1440; '
            'không còn thẻ bo góc (0 tham chiếu `.ops-wallet-card` trong DOM và trong sheet). '
            'AC2 (một phần): hai câu phụ đề chỉ lặp lại nhãn đã bỏ; dòng bút toán sổ quỹ vẫn theo '
            'băng thẻ dùng chung `tt-table` — KHÔNG sửa, xem "chưa làm". AC3: ảnh 390/768/1440.',
            'Sửa thêm (phát hiện khi điều tra flake): `OpsFundBookSection` đọc `data.period.from` '
            'không chặn → payload thiếu `period` ném lỗi unhandled, kéo cả trang (và cả file test) '
            'sập. Trước: 3/6 lần chạy file test đỏ (2 lần đỏ dây chuyền 12 test); sau khi chặn + '
            'sửa mock: 3/3 xanh. `pnpm design:drift` không tăng (hexFallback ban đầu bị ratchet bắt, '
            'đã đổi sang `var(--warning-text)`).',
            'Chưa làm: mật độ DÒNG của sổ quỹ ở 390 (6 sự kiện, mỗi nhãn một dòng) — đây là băng '
            '`tt-table` dùng chung, không phải hình dạng riêng của trang.',
        ],
        [
            'Mở /ops/wallet (giaonhan) ở 390px — ba số dư nằm trên một dải kẻ, không còn thẻ.',
            'Ở 768/1440 dải vẫn một hàng, không tràn ngang.',
        ],
    ),
}


def main():
    for name, (status, blocks, steps) in DONE.items():
        src = None
        for column in ('IN_PROGRESS', 'TODO'):
            candidate = KANBAN / column / name
            if candidate.exists():
                src = candidate
                break
        if src is None:
            print(f'!! not on the board: {name}')
            continue
        doc = docx.Document(str(src))
        for paragraph in doc.paragraphs:
            if paragraph.text.startswith('Trạng thái:'):
                for run in paragraph.runs[1:]:
                    run.text = ''
                if paragraph.runs:
                    paragraph.runs[0].text = status
                else:
                    paragraph.add_run(status)
                break
        for block in blocks:
            doc.add_paragraph(block)
        if steps:
            doc.add_paragraph('Hướng dẫn xác minh:')
            for index, step in enumerate(steps, start=1):
                doc.add_paragraph(f'{index}. {step}')
        dest = KANBAN / 'DEV_COMPLETED' / name
        doc.save(str(dest))
        src.unlink()
        print(f'moved {name}')


if __name__ == '__main__':
    main()
