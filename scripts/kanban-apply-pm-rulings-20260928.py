#!/usr/bin/env python3
"""Record the 2026-09-28 PM rulings inside the kanban cards they unblock.

The board is the team's shared progress signal: a card that still reads
"CẦN THÔNG TIN" when the ruling has landed is lying, and the next lane (or
the next PM question) will be routed to a decision that has already been
made. Each ruling is embedded in the card itself (per the board's
self-contained rule) rather than referenced by a local path.
"""
import sys
from docx import Document

RULINGS = {
    "20260928_167-two-fund-accounts-voucher.docx": {
        "status": "Trạng thái: MỞ — đã có ruling 2026-09-28; phần còn lại là vận hành, không còn hỏi PM",
        "block": [
            "[2026-09-28] RULING PM (ADR 2026-09-28-kanban-pm-open-questions-rulings) — câu hỏi "
            "'Lồng ghép Quỹ vào luôn không?' ĐÃ ĐƯỢC TRẢ LỜI, và phần lớn thẻ đã có sẵn trong code:",
            "• Hai nguồn quỹ GIỮ RIÊNG, tích hợp vào màn phơi phiếu. Không gộp thành một sổ, không thêm "
            "hạn mức/hạn tín dụng (PM không hề yêu cầu, thêm vào là tự bịa phạm vi).",
            "• Sổ quỹ tách hai nguồn đã CHỐT và đã SHIP: PRD OpsVanHanh §5.2 ghi rõ '(chốt 21/09, ship 22/09)' — "
            "mỗi phiếu gắn đúng một nguồn, sổ từng nguồn đọc riêng, tài khoản chưa gắn nguồn đếm riêng.",
            "• Cổng chặn gán sai quỹ ĐÃ CÓ: VOUCHER_REQUIRED_FUND (treasury.service.ts:538-547) ánh xạ mọi nhóm chi "
            "phí sang đúng nguồn PM quy định, assertVoucherFundMatches (treasury.service.ts:562) từ chối phiếu "
            "lệch nguồn và phiếu trộn hai nguồn — commit 26d19d1f, 6 test trong voucher-fund-direction.test.ts "
            "(có test chống trôi: đọc mapping từ source thay vì chép lại).",
            "• Còn lại cho thẻ: đối chiếu nốt các tiêu chí nghiệm thu chưa tick, không phải viết thêm quỹ.",
        ],
    },
    "20260928_169-advance-reimbursement-report.docx": {
        "status": "Trạng thái: MỞ — 3 câu hỏi PM đã có ruling 2026-09-28, tiêu chí 1 + 5 đã chốt lại",
        "block": [
            "[2026-09-28] RULING PM (ADR 20260928) — 3 câu hỏi đã được trả lời từ PRD §9.2:",
            "• Câu 1, 'đợt làm đề nghị' = LÔ ĐỐI SOÁT (expenseReconciliations). PRD định nghĩa 'đợt' bằng "
            "ngữ nghĩa: một đợt SỞ HỮU một tập chi phí, PHÂN BỔ các khoản ứng đã nhận, và không khoản nào "
            "được tính vào hai đợt. expenseReconciliations là cấu trúc duy nhất có đủ cả ba (có opsUserId, "
            "khung from/to, advanceAmount, và expense_reconciliation_advances gán ứng vào từng đợt). "
            "debitSettlementRounds (chốt công nợ tháng) không có các thuộc tính này và thuộc về công nợ nhà xe.",
            "• Câu 2, trường nào cho vào báo cáo: from/to + nhân viên + lô + khách hàng; bốn trường đã có sẵn "
            "trong expenseListQuerySchema (cùng schema với doanh thu nên đã được kiểm thử). Thiếu duy nhất "
            "là trục 'đợt'.",
            "• Câu 3, tiêu chí 5 ĐÃ CHỐT LẠI: câu chữ cũ 'số còn phải hoàn ứng bằng số dư quỹ' SAI và phải sửa. "
            "PRD nói thẳng: 'Số dư quỹ dùng chiều ngược lại … không ép hai số có cùng dấu' và 'sau khi phiếu post, "
            "sổ quỹ và báo cáo hội tụ về một số'. Test đúng là: |remaining| của báo cáo bằng bookBalance của tài "
            "khoản quỹ đó (dấu ngược), và hội tụ sau khi phiếu post — KHÔNG phải bằng số.",
        ],
    },
    "20260928_171-cost-detail-view-edit.docx": {
        "status": "Trạng thái: MỞ — đã có ruling 2026-09-28; cột Người thanh toán đã có sẵn trong code",
        "block": [
            "[2026-09-28] RULING PM (ADR 20260928) — tiêu chí 4 và 6 đã được trả lời, và tiêu chí 6 ĐÃ CÓ CODE:",
            "• Tiêu chí 6, cột 'Người thanh toán': PRD OpsVanHanh §9.1 — \"'Người thanh toán' là người thực hiện "
            "khoản chi; nhập thay không đổi người này thành người đang đăng nhập\". Code đã hiện thực đúng: "
            "phoi-phieu-control.service.ts:472-473 đọc opsExpenseEntries.paidById nối users.fullName, trả về "
            "payerName/payerUserId; PhoiPhieuChiHoDialog.tsx:145,156 render cột 'Người thanh toán'.",
            "• Phần còn thiếu của tiêu chí 6 là mệnh đề thứ hai ('nhập thay không đổi') — cần một test khẳng định "
            "khi kế toán nhập hộ mà dòng đó vẫn giữ người trả thật, không bị ghi đè bằng user đang đăng nhập.",
            "• Tiêu chí 4, dòng lái xe trong chi tiết chi hộ: văn bản PM liệt kê chi tiết chi hộ gồm chi phí "
            "của Ops VÀ của lái xe, nên dòng lái xe thuộc màn này.",
        ],
    },
    "20260928_173-monthly-payable-receivable-reports.docx": {
        "status": "Trạng thái: MỞ — tiêu chí 2 và 3 đã có ruling 2026-09-28 từ văn bản PM",
        "block": [
            "[2026-09-28] RULING PM (ADR 20260928) — tiêu chí 2 và 3 đã có câu trả lời nguyên văn trong tài liệu "
            "'các chi phí' của PM:",
            "• Tiêu chí 2, thứ tự hiển thị: \"Ưu tiên thứ tự, với những khách xhd nhiều lần 1 tháng, được ưu tiên "
            "xếp nối tiếp\" — sắp theo số lượt giao dịch trong kỳ giảm dần để khách lặp lại nằm cạnh nhau, kèm "
            "tiebreaker ổn định (ví dụ tên) để thứ tự không nhảy giữa hai lần render.",
            "• Tiêu chí 3, gộp một dòng: \"Với khách hàng có phát sinh cả thu / trả 1 tháng … Ưu tiên hiển thị tổng "
            "hợp trên cùng 1 dòng: cả cước phải thu / phải trả, số lượng\" — một chủ thể có phát sinh CẢ hai "
            "chiều trong kỳ lọc thì hiển thị MỘT dòng gộp mang cả hai con số và số lượng, không tách thành hai dòng.",
            "• Xe nhà Silver Sea tính như một mã nhà xe thường trong báo cáo phải trả (không loại trừ).",
            "• Hai bảng dùng CHUNG một component render — sửa layout một bảng thì bảng kia đổi theo.",
        ],
    },
}


def main(board: str) -> int:
    import os
    touched = 0
    for name, payload in RULINGS.items():
        for col in ("TODO", "IN_PROGRESS", "DEV_COMPLETED"):
            path = os.path.join(board, col, name)
            if not os.path.exists(path):
                continue
            doc = Document(path)
            hit = False
            for para in doc.paragraphs:
                if para.text.strip().startswith("Trạng thái:"):
                    for run in para.runs[1:]:
                        run.text = ""
                    para.runs[0].text = payload["status"]
                    hit = True
                    break
            if not hit:
                print(f"  !! no Trạng thái: paragraph in {name}", file=sys.stderr)
                return 1
            doc.add_paragraph("")
            doc.add_paragraph(
                f"[2026-09-28] RULING PM — câu hỏi đã chốt, thẻ không còn chờ PM "
                f"(xem ADR 2026-09-28-kanban-pm-open-questions-rulings trong repo)"
            )
            for line in payload["block"]:
                doc.add_paragraph(line)
            doc.save(path)
            print(f"updated {col}/{name}")
            touched += 1
            break
    print(f"touched {touched} card(s)")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
