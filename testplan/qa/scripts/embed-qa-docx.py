#!/usr/bin/env python3
"""Append a QA evidence block (text + embedded images) to a SilverSea kanban card
and rewrite its `Trạng thái:` line. Driven by a JSON spec so the block is
reproducible and no local path ever lands inside the docx.

Spec:
{
  "card": "/abs/path/card.docx",
  "status": "QA PASSED — đã xác minh trên staging",
  "sections": [
    {"heading": "[2026-10-05] QA PASSED — staging: <what was verified>"},
    {"text": "..."},
    {"image": "/abs/path/x.png", "caption": "...", "width": 6.0},
    ...
  ]
}
"""
import json
import sys

from docx import Document
from docx.shared import Inches


def main(spec_path: str) -> None:
    with open(spec_path, encoding="utf-8") as fh:
        spec = json.load(fh)

    doc = Document(spec["card"])

    for section in spec.get("sections", []):
        if "heading" in section:
            doc.add_paragraph(section["heading"])
        if "text" in section:
            for line in section["text"].split("\n"):
                doc.add_paragraph(line)
        if "image" in section:
            if section.get("caption"):
                doc.add_paragraph(section["caption"])
            doc.add_picture(section["image"], width=Inches(section.get("width", 6.0)))

    status = spec["status"]
    targets = [p for p in doc.paragraphs if p.text.strip().startswith("Trạng thái:")]
    hit = False
    if targets:
        para = targets[-1]  # the live status line is the last one; earlier ones are history
        for run in para.runs:
            run.text = ""
        if para.runs:
            para.runs[0].text = f"Trạng thái: {status}"
        else:
            para.add_run(f"Trạng thái: {status}")
        hit = True

    if not hit:
        doc.add_paragraph(f"Trạng thái: {status}")

    doc.save(spec["card"])
    print(f"OK card={spec['card']} status={status!r} sections={len(spec.get('sections', []))}")


if __name__ == "__main__":
    main(sys.argv[1])
